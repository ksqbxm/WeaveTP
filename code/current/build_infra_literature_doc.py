from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("infra_literature_dynamic_tp_migration.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(92, 99, 107)
INK = RGBColor(35, 38, 42)
LINK = RGBColor(5, 99, 193)


def set_run_font(run, name="Calibri", east_asia="SimSun", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Calibri", east_asia="SimSun", size=11, color=INK, bold=None):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "SimSun")
    r_pr.append(r_fonts)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=9, color=MUTED)


def set_table_borders(table, color="D9E2F3", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            cell.vertical_alignment = 1


def add_labeled_paragraph(doc, label, text, style="Paper Meta"):
    p = doc.add_paragraph(style=style)
    label_run = p.add_run(label)
    set_run_font(label_run, size=10.5, color=NAVY, bold=True)
    text_run = p.add_run(text)
    set_run_font(text_run, size=10.5, color=INK)
    return p


def add_source_paragraph(doc, paper_url, code_url=None, code_label="代码仓库"):
    p = doc.add_paragraph(style="Source")
    run = p.add_run("来源：")
    set_run_font(run, size=9.5, color=MUTED, bold=True)
    add_hyperlink(p, "论文页面", paper_url)
    if code_url:
        sep = p.add_run("  |  ")
        set_run_font(sep, size=9.5, color=MUTED)
        add_hyperlink(p, code_label, code_url)
    return p


def add_paper(doc, index, paper):
    heading = doc.add_paragraph(style="Heading 3")
    heading.paragraph_format.keep_with_next = True
    run = heading.add_run(f"{index}. {paper['title']}")
    set_run_font(run, name="Calibri", east_asia="Microsoft YaHei", size=12, color=DARK_BLUE, bold=True)
    add_labeled_paragraph(doc, "年份 / 状态：", paper["year_status"])
    add_labeled_paragraph(doc, "作者：", paper["authors"])
    add_labeled_paragraph(doc, "会议或期刊：", paper["venue"])
    add_labeled_paragraph(doc, "公开代码：", paper["code_status"])
    add_labeled_paragraph(doc, "主要内容：", paper["content"], style="Paper Body")
    add_labeled_paragraph(doc, "与 MOETP++ 的关系：", paper["relation"], style="Paper Body")
    add_source_paragraph(doc, paper["paper_url"], paper.get("code_url"), paper.get("code_label", "代码仓库"))


dynamic_papers = [
    {
        "title": "Llumnix: Dynamic Scheduling for Large Language Model Serving",
        "year_status": "2024；正式发表",
        "authors": "Biao Sun, Ziming Huang, Hanyu Zhao, Wencong Xiao, Xinyi Zhang, Yong Li, Wei Lin",
        "venue": "18th USENIX Symposium on Operating Systems Design and Implementation (OSDI 2024), pp. 173-191",
        "code_status": "公开；官方仓库和 OSDI artifact 均可获得",
        "content": "面向异构且不可预测的在线请求，在多个模型实例之间进行运行时重调度；通过请求及内存状态的 live migration 改善负载均衡、资源碎片和 SLO 隔离。其核心迁移对象是请求状态，而不是 TP 拓扑。",
        "relation": "适合作为请求级动态迁移基线，用来验证 MOETP++ 在动态负载下是否能减少队列和尾延迟；不能直接代替 TP 迁移基线。",
        "paper_url": "https://www.usenix.org/conference/osdi24/presentation/sun-biao",
        "code_url": "https://github.com/llumnix-project/llumnix-ray",
    },
    {
        "title": "ServerlessLLM: Low-Latency Serverless Inference for Large Language Models",
        "year_status": "2024；正式发表",
        "authors": "Yao Fu, Leyang Xue, Yeqi Huang, Andrei-Octavian Brabete, Dmitrii Ustiugov, Yuvraj Patel, Luo Mai",
        "venue": "18th USENIX Symposium on Operating Systems Design and Implementation (OSDI 2024)",
        "code_status": "公开；官方 GitHub 仓库",
        "content": "通过多层次 checkpoint 加载、面向推理的 live migration 和启动时模型调度，利用本地存储与内存层次降低模型实例启动及迁移的延迟。",
        "relation": "可借鉴模型状态迁移、局部性和重算替代大块状态传输的思路，但不直接改变 TP 度数。",
        "paper_url": "https://arxiv.org/abs/2401.14351",
        "code_url": "https://github.com/ServerlessLLM/ServerlessLLM",
    },
    {
        "title": "BlitzScale: Fast and Live Large Model Autoscaling with O(1) Host Caching",
        "year_status": "2025；正式发表",
        "authors": "Dingyan Zhang, Haotian Wang, Yang Liu, Xingda Wei, Yizhou Shan, Rong Chen, Haibo Chen",
        "venue": "19th USENIX Symposium on Operating Systems Design and Implementation (OSDI 2025), pp. 275-293",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "利用 GPU 计算网络和网络优化 multicast 加速参数加载，并将扩缩容粒度从实例级细化到 layer 级，使新实例能够在参数完全加载前承担部分层计算。",
        "relation": "适合吸收在线扩缩容和不中断迁移的设计；可作为 MOETP++ 的资源弹性补充，但不是 TP/KV 迁移算法。",
        "paper_url": "https://www.usenix.org/conference/osdi25/presentation/zhang-dingyan",
    },
    {
        "title": "FuseLink: Enabling Efficient GPU Communication over Multiple NICs",
        "year_status": "2025；正式发表",
        "authors": "Zhenghang Ren, Yuxuan Li, Zilong Wang, Xinyang Huang, Wenxue Li, Kaiqiang Xu, Xudong Liao, Yijun Sun, Bowen Liu, Han Tian, Junxue Zhang, Mingfei Wang, Zhizhen Zhong, Guyue Liu, Ying Zhang, Kai Chen",
        "venue": "19th USENIX Symposium on Operating Systems Design and Implementation (OSDI 2025), pp. 91-108",
        "code_status": "截至检索时未确认有独立公开复现仓库；论文实现集成到 NCCL",
        "content": "针对静态 GPU-NIC 绑定造成的 NIC 热点，引入 GPU 中继和多 NIC 流量转发，让通信流量使用空闲 NIC，改善 LLM serving 和 MoE all-to-all 的动态不均衡。",
        "relation": "最适合吸收到 MOETP++ 的网络层：迁移决策不仅看 GPU，还应考虑每条链路和 NIC 的实时有效带宽。",
        "paper_url": "https://www.usenix.org/conference/osdi25/presentation/ren",
    },
    {
        "title": "Helix: Serving Large Language Models over Heterogeneous GPUs and Network via Max-Flow",
        "year_status": "2025；正式发表",
        "authors": "Yixuan Mei, Yonghao Zhuang, Xupeng Miao, Juncheng Yang, Zhihao Jia, Rashmi Vinayak",
        "venue": "30th ACM International Conference on Architectural Support for Programming Languages and Operating Systems (ASPLOS 2025)",
        "code_status": "公开；官方实现和 artifact",
        "content": "将异构 GPU、节点和网络链路建模为容量约束，用 max-flow/MILP 联合决定模型放置与请求调度，适配不同 GPU 算力和非均匀网络带宽。",
        "relation": "可作为 MOETP++ 的异构带宽调度基线或 planner 组件；Helix 主要做放置和调度，不负责在线 TP/KV 迁移。",
        "paper_url": "https://www.cs.cmu.edu/~rvinayak/papers/Helix_ASPLOS_2025_Serving_LLMs_over_Heterogeneous_GPUs_and_Network_via_Max_Flow.pdf",
        "code_url": "https://github.com/Thesys-lab/Helix-ASPLOS25",
    },
    {
        "title": "ThunderServe: High-performance and Cost-efficient LLM Serving in Cloud Environments",
        "year_status": "2025；正式发表",
        "authors": "Youhe Jiang, Fangcheng Fu, Xiaozhe Yao, Taiyi Wang, Bin Cui, Ana Klimovic, Eiko Yoneki",
        "venue": "Proceedings of Machine Learning and Systems 7 (MLSys 2025)",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "面向异构云 GPU 和网络环境，使用部署规划算法适配不同资源容量，并设计轻量级在线重调度来应对节点故障和工作负载变化，避免重启正在运行的服务。",
        "relation": "可作为动态非均匀资源场景的系统级对照；与 MOETP++ 的区别是它不处理 TP shard 或 KV cache 的细粒度迁移。",
        "paper_url": "https://proceedings.mlsys.org/paper_files/paper/2025/hash/c2a0e26dd9ee7d57e92bb1c24b39659a-Abstract-Conference.html",
    },
    {
        "title": "Dilu: Enabling GPU Resourcing-on-Demand for Serverless DL Serving via Introspective Elasticity",
        "year_status": "2025；正式发表",
        "authors": "Cunchi Lv, Xiao Shi, Zhengyu Lei, Jinyue Huang, Wenting Tan, Xiaohui Zheng, Xiaofang Zhao",
        "venue": "ACM International Conference on Architectural Support for Programming Languages and Operating Systems (ASPLOS 2025)",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "提出 introspective elasticity，通过多因素 profiling、资源互补调度和自适应二维 co-scaling，同时调节单任务资源和任务副本数量，降低动态 serverless GPU 碎片。",
        "relation": "适合借鉴资源弹性和碎片管理；可扩展 MOETP++ 的 GPU 分配层，但不直接解决 TP 迁移。",
        "paper_url": "https://arxiv.org/abs/2503.05130",
    },
    {
        "title": "EcoServe: Efficient LLM Serving on Commodity GPU Clusters with Data-Reduced Cross-Instance Orchestration",
        "year_status": "2026；正式发表",
        "authors": "Jiangsu Du, Hongbin Zhang, Taosheng Wei, Zhenyi Zheng, Jiazhi Jiang, Kaiyi Wu, Zhiguang Chen, Yutong Lu",
        "venue": "20th USENIX Symposium on Operating Systems Design and Implementation (OSDI 2026), pp. 1787-1802",
        "code_status": "公开；官方 GitHub 仓库",
        "content": "面向没有 NVLink/InfiniBand 的普通 GPU 集群，采用部分解耦的 prefill-decode 协同、rolling activation、宏实例调度和 mitosis scaling，在 Ethernet 环境中做在线容量调整。",
        "relation": "与带宽受限的 MOETP++ 实验环境较接近，可作为普通网络集群下的动态调度参照；不改变 TP 拓扑。",
        "paper_url": "https://www.usenix.org/conference/osdi26/presentation/du",
        "code_url": "https://github.com/MLSysU/EcoServe",
    },
    {
        "title": "Connex: Endpoint Mobility Primitives for Dynamic LLM Serving",
        "year_status": "2026；SIGCOMM 2026 正式会议论文",
        "authors": "Yanying Lin, Vincent Liu, Tao Luo, ChengZhong Xu, Kejiang Ye",
        "venue": "ACM SIGCOMM 2026",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "把 worker 的加入、退出和迁移作为通信层的一等原语，通过 epoch-based routing、显式 handover 和 credit-based backpressure，在 token、activation 或 KV 传输进行时维持顺序和流量隔离。",
        "relation": "适合借鉴 TP 迁移期间的通信切换协议；它是通信基础设施，不是 TP 重分片算法。",
        "paper_url": "https://conferences.sigcomm.org/sigcomm/2026/program/papers/",
    },
    {
        "title": "KVServe: Service-Aware KV Cache Compression for Communication-Efficient Disaggregated LLM Serving",
        "year_status": "2026；SIGCOMM 2026 正式会议论文",
        "authors": "Zedong Liu, Xinyang Ma, Dejun Luo, Hairui Zhao, Bing Lu, Wenjing Huang, Yida Gu, Xingchen Liu, Zheng Wei, Jinyang Liu, Dingwen Tao, Guangming Tan",
        "venue": "ACM SIGCOMM 2026",
        "code_status": "公开；官方 GitHub 仓库",
        "content": "将 KV 压缩策略组织成可组合的策略空间，使用 Bayesian profiling 和 service-aware online controller，根据工作负载、网络带宽和 SLO 动态选择压缩配置，并用流水线 fetcher 降低 KV 通信瓶颈。",
        "relation": "可以作为 MOETP++ 的带宽感知迁移代价模型和 KV 传输优化模块；不直接切换 TP。",
        "paper_url": "https://conferences.sigcomm.org/sigcomm/2026/program/papers/",
        "code_url": "https://github.com/hpdps-group/KVServe",
    },
    {
        "title": "Accelerating Distributed MoE Training and Inference with Lina",
        "year_status": "2023；正式发表",
        "authors": "Jiamin Li, Yimin Jiang, Yibo Zhu, Cong Wang, Hong Xu",
        "venue": "2023 USENIX Annual Technical Conference (ATC 2023), pp. 945-959",
        "code_status": "论文公开；截至检索时未确认官方 GitHub 仓库",
        "content": "分析 MoE all-to-all 的训练和推理瓶颈；推理时依据专家流行度预测动态分配设备，使热门专家获得更多资源、冷门专家被压缩到更少设备，从而平衡传输量和链路带宽。",
        "relation": "对 DeepSeek 类 MoE 的专家负载动态非常有参考价值，可与 MOETP++ 的专家/TP 联合调度结合。",
        "paper_url": "https://www.usenix.org/conference/atc23/presentation/li-jiamin",
    },
    {
        "title": "Tutel: Adaptive Mixture-of-Experts at Scale",
        "year_status": "2023；正式发表",
        "authors": "Changho Hwang, Wei Cui, Yifan Xiong, Ziyue Yang, Ze Liu, Han Hu, Zilong Wang, Rafael Salas, Jithin Jose, Prabhat Ram, HoYuen Chau, Peng Cheng, Fan Yang, Mao Yang, Yongqiang Xiong",
        "venue": "Proceedings of Machine Learning and Systems 5 (MLSys 2023)",
        "code_status": "公开；Microsoft 官方仓库",
        "content": "提供动态自适应并行和流水线；通过为参数与输入设计一致布局，实现运行时切换而无需 tensor migration，并配合 Flexible All-to-All、2DH All-to-All 等通信优化。",
        "relation": "最值得吸收的是“布局不变、切换不搬运”的思想；论文实验以 SwinV2-MoE 为主，不能直接等同于 DeepSeek 在线推理。",
        "paper_url": "https://proceedings.mlsys.org/paper_files/paper/2023/hash/5616d34cf8ff73942cfd5aa922842556-Abstract-mlsys2023.html",
        "code_url": "https://github.com/microsoft/tutel",
    },
    {
        "title": "Comet: Fine-grained Computation-communication Overlapping for Mixture-of-Experts",
        "year_status": "2025；正式发表",
        "authors": "Shulai Zhang, Ningxin Zheng, Haibin Lin, Ziheng Jiang, Wenlei Bao, Chengquan Jiang, Qi Hou, Weihao Cui, Size Zheng, Li-Wen Chang, Quan Chen, Xin Liu",
        "venue": "Proceedings of Machine Learning and Systems 7 (MLSys 2025)",
        "code_status": "论文称实现已进入 ByteDance Flux；需按版本核对可复现性",
        "content": "通过数据依赖分析、任务重排和自适应 workload assignment，将 MoE 通信切分到更细粒度并与计算重叠，降低 all-to-all 对端到端执行的阻塞。",
        "relation": "适合隐藏 MOETP++ 的迁移/通信开销，作为迁移执行器的 overlap 优化，而不是独立的 TP 调度器。",
        "paper_url": "https://seed.bytedance.com/en/public_papers/comet-fine-grained-computation-communication-overlapping-for-mixture-of-experts",
        "code_url": "https://github.com/bytedance/flux",
        "code_label": "Flux 代码",
    },
    {
        "title": "Balancing and Beyond: Communication-Centric Optimizations in Expert Parallelism (EPIC)",
        "year_status": "2026；SIGCOMM 2026 正式会议论文",
        "authors": "Jiamin Cao, Qingxu Li, Yaozhong Liu, Jiaqi Gao, Yan Zhang, Shangfeng Shi, Zian Chen, Yizhi Wang, Jun Zhang, Kunling He, Ennan Zhai, Jianbo Dong, Binzhang Fu, Dennis Cai",
        "venue": "ACM SIGCOMM 2026",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "提出 EPIC，渐进式解决生产 Expert Parallelism 的专家负载不均和通信低效：性能感知专家迁移、运行时专家激活、拓扑自适应传输 kernel，以及细粒度计算-通信重叠。",
        "relation": "这是目前与 MOETP++ 的 MoE 动态专家迁移和非均匀通信场景最接近的系统论文之一。",
        "paper_url": "https://conferences.sigcomm.org/sigcomm/2026/program/papers/",
    },
    {
        "title": "CRAFT: Fine-Grained Cost-Aware Expert Replication for Efficient MoE Serving",
        "year_status": "2026；正式发表",
        "authors": "Adrian Zhao, Zhenkun Cai, Zhenyu Song, Lingfan Yu, Haozheng Fan, Jun Wu, Yida Wang, Nandita Vijaykumar",
        "venue": "Proceedings of Machine Learning and Systems 8 (MLSys 2026)",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "在固定显存预算下按照每层、每个专家的预期收益进行细粒度复制，避免传统方案过度复制低收益专家，以提高专家利用率和服务吞吐。",
        "relation": "可作为 MOETP++ 的静态/半动态专家复制对照，帮助区分“复制缓解热点”和“迁移解决动态热点”的收益。",
        "paper_url": "https://proceedings.mlsys.org/paper_files/paper/2026/hash/3a7f9e485845dac27423375c934cb4db-Abstract-Conference.html",
    },
    {
        "title": "Capacity-Aware Inference: Mitigating the Straggler Effect in Mixture of Experts",
        "year_status": "2026；正式发表，但属于机器学习会议而非系统顶会",
        "authors": "Shwai He, Weilin Cai, Jiayi Huang, Ang Li",
        "venue": "International Conference on Learning Representations (ICLR 2026)",
        "code_status": "公开；官方 GitHub 仓库",
        "content": "在推理时限制专家容量并丢弃或重路由溢出 token，利用本地低负载专家缓解最忙专家造成的 straggler effect。",
        "relation": "可作为 MoE 负载均衡的算法级补充基线，但它不处理 TP 拓扑或 KV 迁移。",
        "paper_url": "https://proceedings.iclr.cc/paper_files/paper/2026/hash/94e845868a9ace4bc239d0c529d32f4c-Abstract-Conference.html",
        "code_url": "https://github.com/CASE-Lab-UMD/Capacity-Aware-MoE",
    },
]


tp_papers = [
    {
        "title": "Enabling Parallelism Hot Switching for Efficient Training of Large Language Models (HotSPa)",
        "year_status": "2024；正式发表",
        "authors": "Hao Ge, Fangcheng Fu, Haoyang Li, Xuanyu Wang, Sheng Lin, Yujie Wang, Xiaonan Nie, Hailin Zhang, Xupeng Miao, Bin Cui",
        "venue": "ACM SIGOPS 30th Symposium on Operating Systems Principles (SOSP 2024), pp. 178-194",
        "code_status": "公开；论文作者的 Hetu 仓库",
        "content": "针对序列长度变化导致的最优并行策略变化，将一个 mini-batch 划分为多个组并采用不同策略；用共享模型存储、图编译器和 hot-switch planner 在线传输参数与梯度。",
        "relation": "提供了顶会级的并行策略热切换和通信计划思想，但目标是训练；MOETP++ 需要把 planner 和状态一致性机制改造成推理/KV 场景。",
        "paper_url": "https://sigops.org/s/conferences/sosp/2024/accepted.html",
        "code_url": "https://github.com/PKU-DAIR/Hetu",
    },
    {
        "title": "Seesaw: High-throughput LLM Inference via Model Re-sharding",
        "year_status": "2025；正式发表",
        "authors": "Qidong Su, Wei Zhao, Xin Li, Muralidhar Andoorveedu, Chenhao Jiang, Zhanda Zhu, Kevin Song, Christina Giannoula, Gennady Pekhimenko",
        "venue": "Proceedings of Machine Learning and Systems 7 (MLSys 2025)",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "针对 prefill 与 decode 的计算特征差异，动态重分片模型并使用 tiered KV buffering 和 transition-minimizing scheduling，降低频繁阶段转换造成的开销。",
        "relation": "适合作为重分片和阶段转换的已发表参考；主要面向吞吐型/离线推理，不等同于任意 TP 度数的在线迁移。",
        "paper_url": "https://proceedings.mlsys.org/paper_files/paper/2025/hash/cbc4ab80cd77aa0eb87da062fbcddb46-Abstract-Conference.html",
    },
    {
        "title": "Shift Parallelism: Low-Latency, High-Throughput LLM Inference for Dynamic Workloads",
        "year_status": "2026；ASPLOS 2026 正式会议论文",
        "authors": "Mert Hidayetoglu, Aurick Qiao, Michael Wyatt, Jeff Rasley, Yuxiong He, Samyam Rajbhandari",
        "venue": "ACM International Conference on Architectural Support for Programming Languages and Operating Systems (ASPLOS 2026)",
        "code_status": "公开；Snowflake ArcticInference / vLLM 插件",
        "content": "根据 batch size 在 TP 和 inference-adapted SP 之间切换：低流量使用 TP 降低延迟，高流量使用 SP 提高吞吐。只要 TP x SP 等于总并行度，KV cache layout 保持不变，因此切换不需要大规模 KV 搬运。",
        "relation": "最值得吸收的长处是“通过状态布局不变避免迁移”，可作为 MOETP++ 高频切换的低成本模式。",
        "paper_url": "https://arxiv.org/abs/2509.16495",
        "code_url": "https://github.com/snowflakedb/ArcticInference",
    },
    {
        "title": "FLYING SERVING: On-the-Fly Parallelism Switching for Large Language Model Serving",
        "year_status": "2026；正式发表",
        "authors": "Shouwei Gao, Junqi Yin, Feiyi Wang, Wenqian Dong",
        "venue": "40th ACM International Conference on Supercomputing (ICS 2026), pp. 17-29",
        "code_status": "公开；GitHub 仓库，基于 vLLM",
        "content": "让一组 DP 引擎在运行时合并为更大的 TP group，再按需拆分回 DP；使用 resident weights、zero-copy TP shard view、预初始化 communicator、KV cache adaptor 和 lock-step scheduler。",
        "relation": "与 MOETP++ 的在线 TP/DP 切换目标最接近，建议作为第一优先级复现和对比对象；需要检查其对 DeepSeek-V2-Lite、MoE 和 MLA cache 的适配。",
        "paper_url": "https://arxiv.org/abs/2602.22593",
        "code_url": "https://github.com/shwgao/Flying-Serving",
    },
    {
        "title": "AnchorTP: Resilient LLM Inference with State-Preserving Elastic Tensor Parallelism",
        "year_status": "2026；DATE 2026，属于架构/EDA 方向会议",
        "authors": "Wendong Xu, Chujie Chen, He Xiao, Kuan Li, Jing Xiong, Chen Zhang, Wenyong Zhou, Chaofan Tao, Yang Bai, Bei Yu, Ngai Wong",
        "venue": "IEEE/ACM Design, Automation and Test in Europe Conference (DATE 2026)",
        "code_status": "截至检索时未确认有公开官方仓库",
        "content": "提供非等宽的弹性 TP，使用与推理进程解耦的 daemon 持久化权重和 KV cache，并用 Continuous Minimal Migration 和流水线 P2P/reload 在 GPU 故障后快速恢复。",
        "relation": "适合借鉴 state-preserving elastic TP、最小迁移规划和故障恢复；其主要场景是弹性恢复，不是一般动态带宽调度。",
        "paper_url": "https://arxiv.org/abs/2511.11617",
    },
    {
        "title": "Nitsum: Serving Tiered LLM Requests with Adaptive Tensor Parallelism",
        "year_status": "2026；arXiv 预印本，未确认正式会议发表",
        "authors": "Vikranth Srivatsa, Zijian He, Pu Guo, Dongming Li, Yiying Zhang",
        "venue": "arXiv:2605.05467；预印本",
        "code_status": "截至检索时未确认公开官方仓库",
        "content": "把 TP 度数作为运行时控制变量，与 prefill/decode GPU 划分和分层 SLO 调度联合优化。切换时使用预热进程、TP-aware weight reuse，以及先聚合碎片 KV 再用双缓冲流水线发送的迁移 kernel。",
        "relation": "与 MOETP++ 的动态 TP、带宽感知和 KV 迁移最接近，适合做迁移微基准和算法设计参考；论文状态必须标注为预印本。",
        "paper_url": "https://arxiv.org/abs/2605.05467",
        "code_url": "https://mlsys.wuklab.io/posts/nitsum/",
        "code_label": "作者技术说明",
    },
    {
        "title": "ReMP: Low-Downtime Runtime Model-Parallelism Reconfiguration for LLM Serving",
        "year_status": "2026；arXiv 预印本，未确认正式会议发表",
        "authors": "Haipeng Yuan, Kaining Zheng, Yongshu Bai, Yuchen Zhang, Yunquan Zhang, Baodong Wu, Xiang Gao, Daning Cheng",
        "venue": "arXiv:2606.18741；预印本",
        "code_status": "截至检索时未确认公开官方仓库",
        "content": "把 TP/PP 拓扑与运行时状态解耦，使用共享 CPU 权重存储、预构建并行状态和 standby/wakeup worker；通过二维 KV migration 同时沿 layer 维和 KV-head 维重映射，支持在线拓扑切换。",
        "relation": "是目前最完整的 TP+PP 在线重配置设计之一，可用于校验 MOETP++ 的二维状态映射和迁移事务；但不能按正式顶会论文宣传。",
        "paper_url": "https://arxiv.org/abs/2606.18741",
    },
]


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, size=11, color=INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
]:
    style = styles[style_name]
    set_style_font(style, east_asia="Microsoft YaHei", size=size, color=color, bold=True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.keep_with_next = True

for name, base, size, color in [
    ("Paper Meta", "Normal", 10.5, INK),
    ("Paper Body", "Normal", 10.5, INK),
    ("Source", "Normal", 9.5, MUTED),
]:
    if name not in styles:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles[name]
    style.base_style = styles[base]
    set_style_font(style, size=size, color=color)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(3 if name != "Source" else 9)
    style.paragraph_format.line_spacing = 1.18

# Quiet running furniture for a multi-page reference guide.
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
header_run = header.add_run("AI Infra Literature Guide  |  Dynamic Serving and TP Migration")
set_run_font(header_run, size=9, color=MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_run = footer.add_run("Page ")
set_run_font(footer_run, size=9, color=MUTED)
add_page_field(footer)

# Title block.
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(8)
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(4)
title_run = title.add_run("AI Infra 中动态变化与 TP 迁移论文清单")
set_run_font(title_run, east_asia="Microsoft YaHei", size=24, color=NAVY, bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
subtitle_run = subtitle.add_run("面向 MoE/LLM 推理、动态带宽、资源弹性与运行时模型并行重配置")
set_run_font(subtitle_run, east_asia="Microsoft YaHei", size=12.5, color=MUTED)

meta_table = doc.add_table(rows=3, cols=2)
meta_table.autofit = False
set_table_geometry(meta_table, [1800, 7560], indent_dxa=120)
set_table_borders(meta_table, color="D9E2F3", size="6")
metadata = [
    ("整理日期", "2026-08-26"),
    ("范围", "系统顶会优先：OSDI、SOSP、ASPLOS、SIGCOMM、EuroSys、MLSys、ATC、ICS；另列相关预印本"),
    ("阅读目的", "为 MOETP++ 在动态非均匀带宽、MoE 负载变化和 TP 迁移场景下选择公平对比对象"),
]
for row, (label, value) in zip(meta_table.rows, metadata):
    set_cell_shading(row.cells[0], "E8EEF5")
    p0 = row.cells[0].paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(label)
    set_run_font(r0, size=10, color=NAVY, bold=True)
    p1 = row.cells[1].paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(value)
    set_run_font(r1, size=10, color=INK)
mark_header_row(meta_table.rows[0])

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(12)
note.paragraph_format.space_after = Pt(10)
note.paragraph_format.left_indent = Inches(0.08)
set_paragraph_shading(note, "F4F6F9")
nr = note.add_run("使用说明：")
set_run_font(nr, size=10.5, color=NAVY, bold=True)
nt = note.add_run("请求迁移、PP 重配置、专家迁移和真正的 TP 度数切换不是同一个问题。文档在每篇论文的“与 MOETP++ 的关系”中标出适用范围，避免把不同粒度的方法直接当成同一类基线。")
set_run_font(nt, size=10.5, color=INK)

doc.add_paragraph("快速结论", style="Heading 1")
quick = doc.add_paragraph(style="Normal")
quick.paragraph_format.space_after = Pt(8)
q1 = quick.add_run("优先阅读与对比：")
set_run_font(q1, size=11, color=NAVY, bold=True)
q2 = quick.add_run("Flying Serving（在线 DP/TP 切换）、Shift Parallelism（KV 布局不变）、Helix（异构 GPU/带宽调度）、EPIC/Lina（MoE 专家动态负载）和 Llumnix（请求级动态迁移）。Nitsum 与 ReMP 技术上最接近 TP 迁移，但目前应按预印本处理。")
set_run_font(q2, size=11, color=INK)

doc.add_paragraph("第一部分：动态变化与基础设施调度", style="Heading 1")
intro1 = doc.add_paragraph(style="Normal")
intro1.add_run("本部分覆盖工作负载变化、节点或 GPU 资源变化、网络带宽波动、MoE 专家热度变化以及动态扩缩容。重点观察调度器如何感知变化、如何迁移请求或专家，以及如何利用带宽和通信重叠降低代价。")
for run in intro1.runs:
    set_run_font(run, size=11, color=INK)

doc.add_paragraph("A. LLM serving 与异构资源动态", style="Heading 2")
for idx, paper in enumerate(dynamic_papers[:10], start=1):
    add_paper(doc, idx, paper)

doc.add_paragraph("B. MoE 负载与通信动态补充", style="Heading 2")
for idx, paper in enumerate(dynamic_papers[10:], start=11):
    add_paper(doc, idx, paper)

doc.add_paragraph("第二部分：TP 重配置与迁移", style="Heading 1")
intro2 = doc.add_paragraph(style="Normal")
intro2.add_run("本部分聚焦并行拓扑改变时的权重、通信组、KV cache 和 worker 状态处理。严格意义上的在线 TP 迁移主要见于 Flying Serving、Nitsum、ReMP 和 AnchorTP；Shift、Seesaw、HotSPa 分别代表避免迁移、阶段重分片和训练热切换等相邻路线。")
for run in intro2.runs:
    set_run_font(run, size=11, color=INK)

doc.add_paragraph("TP 重配置与迁移论文", style="Heading 2")
for idx, paper in enumerate(tp_papers, start=1):
    add_paper(doc, idx, paper)

doc.add_paragraph("面向 MOETP++ 的使用建议", style="Heading 1")
recommendations = [
    ("主对比组", "Flying Serving + Shift Parallelism：分别代表在线 TP 切换和通过 KV 布局不变避免迁移。"),
    ("MoE 动态组", "EPIC/Lina + Tutel：分别观察专家迁移/激活、专家热度感知和无迁移自适应布局。"),
    ("网络动态组", "Helix + ThunderServe/FuseLink：构造非均匀带宽、NIC 热点、节点故障和工作负载突变。"),
    ("迁移微基准组", "Nitsum + ReMP + AnchorTP：比较切换时间、KV 迁移字节数、迁移期间吞吐下降和恢复时间，并明确标注预印本或非核心系统会议。"),
    ("统一指标", "switch latency、KV bytes moved、TTFT P50/P99、TPOT、goodput、GPU/link utilization、expert imbalance 和迁移期间中断时间。"),
]
for label, text in recommendations:
    p = doc.add_paragraph(style="Paper Body")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(label + "：")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)

doc.add_paragraph("状态说明", style="Heading 1")
status = doc.add_paragraph(style="Normal")
status.add_run("“正式发表”表示已在会议或期刊正式记录中；“预印本”表示当前只确认 arXiv 版本，不能在论文中写成已被顶会接收。代码状态按 2026-08-26 检索结果记录，公开仓库的版本、依赖和模型支持仍需在实际复现实验前单独核对。")
for run in status.runs:
    set_run_font(run, size=10.5, color=INK)

doc.core_properties.title = "AI Infra 中动态变化与 TP 迁移论文清单"
doc.core_properties.subject = "动态 LLM serving、MoE 负载均衡与运行时 TP 重配置"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.comments = "Literature list prepared from official conference pages, paper pages, and public repositories."

doc.save(OUT)
print(OUT.resolve())
